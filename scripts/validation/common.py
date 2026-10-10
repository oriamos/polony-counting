"""Shared helpers for the real-data validation scripts (docs/app/REAL_DATA_VALIDATION.md).

All input files are untrusted: run every script with ``python -I`` and pass paths as arguments.
Nothing here is imported by the app.
"""
from __future__ import annotations

import csv
import re
from pathlib import Path

import numpy as np
import tifffile

# ---- the counting grid as stored in the ImageJ Overlay of every raw slide -------------
GRID_X0 = 158.0
GRID_Y0 = 158.0
GRID_STEP = 316.23        # px (Overlay lines at 158, 474.23, 790.46, ...)
GRID_COLS = 4
GRID_ROWS = 7


def cell_of(x, y):
    """(row, col) of the 4x7 overlay grid cell containing (x, y); -1 if outside."""
    c = np.floor((np.asarray(x, float) - GRID_X0) / GRID_STEP).astype(int)
    r = np.floor((np.asarray(y, float) - GRID_Y0) / GRID_STEP).astype(int)
    out = (c < 0) | (c >= GRID_COLS) | (r < 0) | (r >= GRID_ROWS)
    return np.where(out, -1, r), np.where(out, -1, c)


def cell_rect(r, c):
    x0 = GRID_X0 + c * GRID_STEP
    y0 = GRID_Y0 + r * GRID_STEP
    return x0, y0, x0 + GRID_STEP, y0 + GRID_STEP


def read_raw(path):
    with tifffile.TiffFile(str(path)) as t:
        arr = t.asarray()
        md = t.imagej_metadata or {}
    return arr, md


def read_points(path):
    """Embedded ImageJ multi-point ROI: (xy[N,2], counter[N], position[N])."""
    import roifile
    with tifffile.TiffFile(str(path)) as t:
        md = t.imagej_metadata or {}
    if not md.get("ROI"):
        return np.zeros((0, 2)), np.zeros(0, int), np.zeros(0, int)
    r = roifile.ImagejRoi.frombytes(md["ROI"])
    xy = np.asarray(r.coordinates(), float)
    cnt = np.asarray(r.counters if r.counters is not None else np.zeros(len(xy)), int)
    pos = np.asarray(r.counter_positions if r.counter_positions is not None else np.zeros(len(xy)), int)
    return xy, cnt, pos


# ---- researcher workbooks --------------------------------------------------------------
def _letters_to_idx(s):
    n = 0
    for ch in s:
        n = n * 26 + ord(ch) - 64
    return n


FIELD_COL0 = _letters_to_idx("H")   # Field-1
N_FIELDS = 24


def counts_sheets(xlsx):
    """{sheet: {(sample, letter): [24 values or None]}} from Counts_*.xlsx.
    Layout: header row with D=letter and 'Field-1' in H, the next row holds the values; the
    sample name is in column C of the 'B' header row (the 'A' row carries a date there)."""
    import openpyxl
    wb = openpyxl.load_workbook(str(xlsx), data_only=True)
    out = {}
    for ws in wb.worksheets:
        rows = {}
        pending_a = None
        for r in range(1, ws.max_row + 1):
            h = ws.cell(r, FIELD_COL0).value
            if not (isinstance(h, str) and h.strip().lower() == "field-1"):
                continue
            letter = ws.cell(r, 4).value
            name = ws.cell(r, 3).value
            vals = [ws.cell(r + 1, FIELD_COL0 + k).value for k in range(N_FIELDS)]
            vals = [None if v is None or v == "" else v for v in vals]
            if letter == "A":
                pending_a = (r, name, vals)
            elif letter == "B":
                sample = str(name).strip() if name else "?"
                rows[(sample, "B")] = vals
                if pending_a is not None:
                    a_name = pending_a[1]
                    # the A header carries a date, or (Controls sheet) the sample name
                    if a_name and not re.match(r"\d{2}-\d{2}-\d{4}", str(a_name)):
                        sample_a = str(a_name).strip()
                        rows.pop((sample, "B"))
                        sample = sample_a if not name else sample
                        rows[(sample, "B")] = vals
                    rows[(sample, "A")] = pending_a[2]
                pending_a = None
        if pending_a is not None:
            rows[(str(pending_a[1]), "A")] = pending_a[2]
        out[ws.title] = rows
    return out


def t7_csv(path):
    """{slide_id: [24 values or None]} from a T7A/T7B CSV."""
    out = {}
    with open(path, newline="", encoding="utf-8-sig") as f:
        rd = csv.reader(f)
        header = next(rd)
        fi = [i for i, h in enumerate(header) if re.fullmatch(r"field \d+", h.strip().lower())]
        for row in rd:
            if not row or not row[0].strip():
                continue
            vals = []
            for i in fi:
                v = row[i].strip() if i < len(row) else ""
                vals.append(None if v == "" else float(v))
            out[row[0].strip()] = vals
    return out


def field_vector(vals):
    return [v for v in vals if v is not None]


def slide_name(path):
    m = re.search(r"_T7_AB_([A-Z]+)-([A-Z])(?:_([A-Z]))?", Path(path).name)
    return (m.group(1), m.group(2), m.group(3)) if m else (Path(path).stem, "?", None)
