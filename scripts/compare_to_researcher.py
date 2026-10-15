#!/usr/bin/env python3
"""Headless comparison of the Polony Review detector with the researcher's embedded ImageJ marks.

For every TIFF in a folder: create an analysis, set the channel mapping, import ONLY the fields the researcher
marked (grid cells that carry her marks), run detection with the base (not fine-tuned) engine, and compare per
field. A field counts only if she provably examined it on that plane (her marks on that plane, or an
examined-fields manifest). Fields without evidence are listed as "unknown" and enter no total or metric.
No UI and no browser: the app is driven in-process through its own HTTP API (FastAPI TestClient).

Output: one Excel file (Summary, Fields, Files, README), rewritten after EVERY file, so a stopped run keeps its results. Everything is DRAFT: agreement with one annotator,
not accuracy.

Usage (from a checkout of the repository, with `pip install -r requirements-app.txt`):

    python compare_to_researcher.py --input /path/to/tiffs --output comparison.xlsx --confirm-page0-red-page1-green

Options: --engine bigfish|hybrid (hybrid needs the AI extra), --manifest examined_fields.json|csv (applied to every
file that it names), --recursive, --data-dir DIR (default: temporary folder, deleted at the end), --limit N.
"""
from __future__ import annotations

import argparse
import shutil
import sys
import tempfile
import time
from pathlib import Path

HERE = Path(__file__).resolve().parent
for cand in (HERE / "src", HERE.parent / "src", HERE.parent.parent / "src", Path.cwd() / "src"):
    if (cand / "polony_app").is_dir():
        sys.path.insert(0, str(cand))
        break

# page (TIFF page = channel index) -> app count types compared with the marks of that page
PLAN = [(0, ["red_positive", "yellow"]), (1, ["green_positive", "yellow"])]

SUMMARY_COLS = ["file", "reference_plane", "app_count_type", "eligible_examined_fields", "positive_marked_fields",
                "verified_zero_fields", "unknown_fields_excluded", "conflict_fields_excluded", "eligibility_source",
                "total_reference", "total_app", "delta_total", "pearson_r", "slope", "mae"]
FIELD_COLS = ["file", "reference_plane", "app_count_type", "field", "region", "status", "examined", "reference_count",
              "app_count", "delta", "reason"]
FILE_COLS = ["file", "state", "reference", "imported_fields", "detection_s", "notes", "error"]

README = [
    "DRAFT. Agreement of the app with ONE annotator's marks, not accuracy: her marks may be incomplete.",
    "Engine = base model (no fine-tuned model is wired into the app); 'before training'.",
    "Only fields the researcher examined are compared: a field counts on a plane only if her marks (or an examined-fields "
    "manifest entry) exist for THAT plane. A field without evidence is 'unknown' and is excluded from totals and metrics; "
    "blank metrics mean 'not defined', never 0.",
    "Plane 0 = red, plane 1 = green, confirmed once by the operator for the whole run (--confirm-page0-red-page1-green).",
    "Compare per field, not slide totals: the researcher counted only her fields.",
    "Files without embedded ImageJ marks are listed in 'Files' as 'no marks'.",
]


def write_workbook(openpyxl, path, summary, fields, filerows):
    """Write the whole workbook to a temp file, then replace the target (a crash never leaves a half-written file)."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Summary"
    for title, cols, data in (("Summary", SUMMARY_COLS, summary), ("Fields", FIELD_COLS, fields),
                              ("Files", FILE_COLS, [[r[k] if k != "notes" else " | ".join(r[k]) for k in FILE_COLS] for r in filerows])):
        sh = ws if title == "Summary" else wb.create_sheet(title)
        sh.append(cols)
        for d in data:
            sh.append(d)
        sh.freeze_panes = "A2"
        for col in sh.columns:
            sh.column_dimensions[col[0].column_letter].width = min(48, max(10, max(len(str(x.value or "")) for x in col[:50]) + 2))
    rd = wb.create_sheet("README")
    for line in README:
        rd.append([line])
    rd.column_dimensions["A"].width = 140
    path = Path(path)
    tmp_path = Path(str(path) + ".part")
    wb.save(tmp_path)
    # On Windows the target can be locked for a moment (Excel open, OneDrive sync, antivirus): retry, never crash the run.
    for attempt in range(6):
        try:
            tmp_path.replace(path)
            return
        except PermissionError:
            time.sleep(1 + attempt)
    alt = path.with_name(f"{path.stem}_{time.strftime('%H%M%S')}{path.suffix}")
    tmp_path.replace(alt)
    print(f"   WARNING: {path.name} is locked (open in Excel?); wrote {alt.name} instead", file=sys.stderr)


def main() -> int:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--input", default=".", help="folder with the TIFF files (default: current folder)")
    ap.add_argument("--output", default="comparison.xlsx", help="Excel file to write (.xlsx), rewritten after every file")
    ap.add_argument("--engine", choices=["bigfish", "hybrid"], default="bigfish")
    ap.add_argument("--manifest", help="examined-fields manifest (JSON or CSV) for files with counted-zero fields")
    ap.add_argument("--recursive", action="store_true")
    ap.add_argument("--limit", type=int, default=0, help="only the first N files (testing)")
    ap.add_argument("--data-dir", help="working directory for the app data (default: temporary)")
    ap.add_argument("--timeout", type=float, default=1800, help="seconds allowed per file for detection")
    ap.add_argument("--confirm-page0-red-page1-green", action="store_true",
                    help="REQUIRED: you confirm that TIFF page 0 is red and page 1 is green for every file")
    a = ap.parse_args()

    if not a.confirm_page0_red_page1_green:
        ap.error("channel colours are never inferred: pass --confirm-page0-red-page1-green after checking them")
    src = Path(a.input)
    if not src.is_dir():
        ap.error(f"not a folder: {src}")
    pat = "**/*" if a.recursive else "*"
    # skip macOS AppleDouble sidecars ("._name.tif": resource-fork stubs copied from a Mac, not images) and __MACOSX folders
    files = sorted(p for p in src.glob(pat) if p.suffix.lower() in (".tif", ".tiff") and p.is_file()
                   and not p.name.startswith("._") and "__MACOSX" not in p.parts)
    if a.limit:
        files = files[: a.limit]
    if not files:
        ap.error("no .tif/.tiff files found")

    try:
        from fastapi.testclient import TestClient
        import openpyxl
        from polony_app.api.main import create_app
    except ImportError as e:  # pragma: no cover
        print(f"missing dependency or wrong location: {e}\n"
              "Run inside the repo checkout after "
              "`pip install -r requirements-app.txt`.", file=sys.stderr)
        return 2

    tmp = None
    data_dir = Path(a.data_dir) if a.data_dir else Path(tempfile.mkdtemp(prefix="polony_cmp_"))
    if not a.data_dir:
        tmp = data_dir
    data_dir.mkdir(parents=True, exist_ok=True)

    summary, fields, filerows = [], [], []
    req = {"engine": a.engine, "preview_window": None}
    if a.engine == "hybrid":
        req["spotiflow"] = {"downscale": 4, "normalizer": "valid_unsaturated"}

    with TestClient(create_app(data_dir)) as c:
        def call(method, url, **kw):
            r = c.request(method, "/api" + url, **kw)
            if r.status_code >= 400:
                try:
                    msg = r.json().get("message") or r.text
                except Exception:
                    msg = r.text
                raise RuntimeError(f"{method} {url}: HTTP {r.status_code} {str(msg)[:200]}")
            return r.json() if r.content else None

        for i, f in enumerate(files, 1):
            row = {"file": f.name, "state": "", "reference": "", "imported_fields": "", "detection_s": "", "notes": [], "error": ""}
            print(f"[{i}/{len(files)}] {f.name}", flush=True)
            try:
                with open(f, "rb") as fh:
                    an = call("POST", "/analyses", files=[("files", (f.name, fh, "image/tiff"))], data={"name": f.stem})
                aid = an["id"]
                call("PUT", f"/analyses/{aid}/channels", json={"mapping": [{"index": 0, "role": "red"}, {"index": 1, "role": "green"}]})
                ref = call("GET", f"/analyses/{aid}/reference")
                if not (ref.get("available") and ref.get("grid")):
                    row["reference"], row["state"] = "no marks", "skipped"
                    row["notes"].append(ref.get("reason") or "no embedded ImageJ marks / grid")
                    filerows.append(row)
                    write_workbook(openpyxl, a.output, summary, fields, filerows)
                    continue
                row["notes"] += ref.get("warnings", [])
                manifest_used = False
                if a.manifest:
                    try:
                        with open(a.manifest, "rb") as mf:
                            m = call("POST", f"/analyses/{aid}/examined-fields",
                                     files=[("file", (Path(a.manifest).name, mf))])
                        manifest_used = m.get("n_entries", 0) > 0
                        row["notes"] += m.get("warnings", [])
                    except RuntimeError as e:
                        row["notes"].append(f"manifest: {e}")
                fi = call("POST", f"/analyses/{aid}/fields/import",
                          json={"mode": "all" if manifest_used else "counted", "shift": [0, 0], "plane": None, "replace": False})
                row["reference"] = "fields imported"
                row["imported_fields"] = sum(1 for r in fi.get("regions", []) if r.get("field_index") is not None)
                t0 = time.time()
                job = call("POST", f"/analyses/{aid}/detect", json=req)
                while job["status"] in ("queued", "running"):
                    if time.time() - t0 > a.timeout:
                        call("POST", f"/jobs/{job['id']}/cancel")
                        raise RuntimeError("detection timed out")
                    time.sleep(0.5)
                    job = call("GET", f"/jobs/{job['id']}")
                if job["status"] != "done":
                    raise RuntimeError(job.get("error") or f"detection {job['status']}")
                row["detection_s"] = round(time.time() - t0, 1)
                for _ in range(20):  # status leaves 'detecting' a moment after the job finishes
                    if call("GET", f"/analyses/{aid}").get("status") != "detecting":
                        break
                    time.sleep(0.5)
                for plane, types in PLAN:
                    for t in types:
                        try:
                            cmp = call("GET", f"/analyses/{aid}/fields/compare", params={"plane": plane, "count_type": t})
                        except RuntimeError as e:
                            row["notes"].append(f"compare plane {plane} {t}: {e}")
                            continue
                        s = cmp["summary"]
                        tr, ta = s.get("total_reference"), s.get("total_app")
                        summary.append([f.name, plane, t, s.get("eligible_examined_fields", 0), s.get("positive_marked_fields", 0),
                                        s.get("verified_zero_fields", 0), s.get("unknown_fields_excluded", 0),
                                        s.get("conflict_fields_excluded", 0), s.get("eligibility_source", "none"), tr, ta,
                                        (ta - tr) if tr is not None and ta is not None else None,
                                        s.get("pearson_r"), s.get("slope"), s.get("mae")])
                        for r in cmp["rows"]:
                            fields.append([f.name, plane, t, r.get("field"), r.get("region"), r.get("status"), r.get("examined"),
                                           r.get("reference_count"), r.get("app_count"), r.get("delta"), r.get("reason")])
                row["state"] = "done"
            except Exception as e:  # one bad file never stops the others
                row["state"], row["error"] = "failed", str(e)
                print(f"   FAILED: {e}", file=sys.stderr)
            filerows.append(row)
            write_workbook(openpyxl, a.output, summary, fields, filerows)

    if tmp and not a.data_dir:
        shutil.rmtree(tmp, ignore_errors=True)
    ok = sum(r["state"] == "done" for r in filerows)
    print(f"done: {ok} compared, {sum(r['state']=='skipped' for r in filerows)} without marks, "
          f"{sum(r['state']=='failed' for r in filerows)} failed -> {a.output}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
