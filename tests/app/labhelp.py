"""Helpers for the lab-export tests: build finalized analyses with grid-field regions, a tiny
Excel-formula evaluator (AVERAGE/SUM/ABS/+-*/ and cell refs) and an optional LibreOffice recalc."""
from __future__ import annotations

import csv
import io
import re
import shutil
import subprocess
import tempfile
from pathlib import Path

import openpyxl

from fakes import counting, exclusion, make_state, obj, rect
from polony_app.schemas import LabCountType, LabSeries, ObjectClass, SlideMeta

TEMPLATES = Path(__file__).resolve().parents[1] / "test_data" / "lab_templates"
# SYNTHETIC files with the layout of the lab's spreadsheets (written by lab_templates/make_templates.py)
T7_CSV = {"T7A": "T7A_template.csv", "T7B": "T7B_template.csv"}
T7_XLSX = {"T7A": "T7A_template.xlsx", "T7B": "T7B_template.xlsx"}
COUNTS_XLSX = "Counts_template.xlsx"
import itertools

_IDS = itertools.count(1)
CELL = 15  # field cell size in px (grid fields on the 160x120 fake image)


def field_region(idx, rid=None):
    """Field `idx` (1-based) = 15x15 px cell on an 8-column grid (first 24 fit in 160x120)."""
    i = idx - 1
    x0, y0 = 5 + (i % 8) * 18, 5 + (i // 8) * 18
    r = counting(rid or f"f{idx}", rect(x0, y0, x0 + CELL, y0 + CELL), name=f"field {idx}")
    r.field_index = idx
    return r, x0, y0


def lab_state(fields, *, slide_id="AI_A", series=(("T7B", LabCountType.yellow),), cls=ObjectClass.yellow,
              sampling="15/09/2026", polony="29/09/2026", volume=2.5, field_area=2.9, grid_area=0.1,
              cf=1, qc=1, finalized=True, extra_regions=(), efficiency=None, aid=None, confirmed=True):
    """`fields`: dict {field_index: n_objects} (a key with 0 = region with zero objects) or a list
    (field 1.. in order; None = no region). Returns a (finalized) AnalysisState."""
    if isinstance(fields, (list, tuple)):
        fields = {i + 1: n for i, n in enumerate(fields) if n is not None}
    regs, objs = [], []
    for idx, n in fields.items():
        r, x0, y0 = field_region(idx)
        regs.append(r)
        for k in range(n):
            objs.append(obj(f"o{idx}_{k}", x0 + 0.5 + (k % CELL), y0 + 0.5 + (k // CELL), cls))
    st = make_state(objs, list(regs) + list(extra_regions))
    a = st.analysis
    a.id = aid or f"{next(_IDS):012x}"
    a.name = slide_id or "slide"
    a.slide = SlideMeta(
        slide_id=slide_id, sampling_date=sampling, polony_date=polony, qc=qc, field_area=field_area,
        grid_area=grid_area, sample_volume_ul=volume, concentration_factor=cf,
        lab_series=[LabSeries(name=n, count_type=t) for n, t in series],
        concentration_formula_confirmed=confirmed, efficiency_input=efficiency)
    if finalized:
        a.finalized_at = "2026-10-01T10:00:00Z"
    return st


# ---------------------------------------------------------------- template readers
def read_template_csv(name):
    raw = (TEMPLATES / name).read_bytes()
    rows = list(csv.reader(io.StringIO(raw.decode("utf-8-sig"), newline="")))
    return raw, rows[0], rows[1:]


def template_slides(name, series, cls):
    """One finalized analysis per row of a template CSV (counts become that many objects)."""
    _, _, rows = read_template_csv(name)
    out = []
    for n, r in enumerate(rows):
        fields = {i + 1: int(v) for i, v in enumerate(r[9:33]) if v != ""}
        ctype = {ObjectClass.red: LabCountType.red_only, ObjectClass.green: LabCountType.green_only,
                 ObjectClass.yellow: LabCountType.yellow}[cls]
        st = lab_state(fields, slide_id=r[0], series=((series, ctype),), cls=cls, sampling=r[1], polony=r[2],
                       cf=float(r[3]), volume=float(r[4]), field_area=float(r[5]), grid_area=float(r[6]),
                       qc=int(r[7]), aid=f"{n + 1:012d}")
        out.append(st.analysis)
    return out


# ---------------------------------------------------------------- formula evaluator
_REF = r"\$?[A-Z]{1,3}\$?\d+"
_TOK = re.compile(rf"(?P<rng>{_REF}:{_REF})|(?P<ref>(?<![A-Za-z_]){_REF}(?![\w(]))")


def _split(ref):
    m = re.match(r"\$?([A-Z]+)\$?(\d+)$", ref)
    return openpyxl.utils.column_index_from_string(m.group(1)), int(m.group(2))


def evaluate(ws, formula, cache=None):
    """Evaluate a (supported) Excel formula on worksheet `ws` (values or nested formulas)."""
    cache = {} if cache is None else cache

    def cell(coord):
        coord = coord.replace("$", "")
        if coord in cache:
            return cache[coord]
        v = ws[coord].value
        if isinstance(v, str) and v.startswith("="):
            v = evaluate(ws, v, cache)
        cache[coord] = v
        return v

    def rng(a, b):
        (c1, r1), (c2, r2) = _split(a), _split(b)
        return [cell(f"{openpyxl.utils.get_column_letter(c)}{r}") for r in range(r1, r2 + 1) for c in range(c1, c2 + 1)]

    def flat(args):
        for x in args:
            if isinstance(x, list):
                yield from (y for y in x if y is not None)
            elif x is not None:
                yield x

    def AVERAGE(*a):
        v = list(flat(a))
        return sum(v) / len(v)

    def SUM(*a):
        return sum(flat(a))

    expr = formula.lstrip("=")
    assert "#REF!" not in expr, formula
    expr = _TOK.sub(lambda m: f"_r({m.group('rng').split(':')[0]!r},{m.group('rng').split(':')[1]!r})" if m.group("rng")
                    else f"_c({m.group('ref')!r})", expr)
    env = {"_r": rng, "_c": lambda c: (cell(c) if cell(c) is not None else 0), "AVERAGE": AVERAGE, "SUM": SUM, "ABS": abs}
    return eval(expr, {"__builtins__": {}}, env)  # noqa: S307 - test-only, formulas written by the exporter


def formula_cells(ws):
    return {c.coordinate: c.value for row in ws.iter_rows() for c in row if isinstance(c.value, str) and c.value.startswith("=")}


# ---------------------------------------------------------------- libreoffice
def soffice():
    return shutil.which("soffice") or shutil.which("libreoffice")


def lo_recalc(xlsx_bytes: bytes):
    """Convert with LibreOffice (recalculates on load) and return the data_only workbook."""
    with tempfile.TemporaryDirectory() as d:
        src = Path(d) / "in.xlsx"
        out = Path(d) / "out"
        src.write_bytes(xlsx_bytes)
        subprocess.run([soffice(), "--headless", f"-env:UserInstallation=file://{d}/profile", "--convert-to", "xlsx",
                        "--outdir", str(out), str(src)], check=True, capture_output=True, timeout=180)
        return openpyxl.load_workbook(out / "in.xlsx", data_only=True)
