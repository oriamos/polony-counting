"""Lab export through the HTTP API (fake detection), incl. field_index round trip and finalize gating."""
import io
import json
import zipfile

import openpyxl
import pytest
from fastapi.testclient import TestClient

from fakes import counting, detector, obj, rect, write_two_channel_tiff  # noqa: F401
from labhelp import T7_CSV, T7_XLSX, TEMPLATES, field_region, read_template_csv
from polony_app.api.main import create_app
from polony_app.schemas import ObjectClass


@pytest.fixture
def client(tmp_path):
    with TestClient(create_app(tmp_path / "data")) as c:
        yield c


@pytest.fixture
def tiff(tmp_path):
    p = tmp_path / "scan.tif"
    write_two_channel_tiff(p)
    return p


def act(client, aid, type_, **payload):
    return client.post(f"/api/analyses/{aid}/actions", json={"type": type_, "payload": payload})


def meta(slide_id, series, **kw):
    m = dict(slide_id=slide_id, sampling_date="15/09/2026", polony_date="29/09/2026", qc=1, field_area=2.9,
             grid_area=0.1, sample_volume_ul=2.5, concentration_factor=1, concentration_formula_confirmed=True,
             lab_series=[{"name": n, "count_type": t, **({"sheet_title": st} if st else {})} for n, t, st in series])
    m.update(kw)
    return m


SERIES = [("T7A", "red_only", "Clade A"), ("T7B", "yellow", "Clade B")]


def make(client, tiff, detector, fields, slide_id="AI_A", series=SERIES, finalize=True, **kw):
    """Upload, map channels, put field regions, fake-detect, approve, set lab metadata, finalize."""
    with open(tiff, "rb") as f:
        aid = client.post("/api/analyses", files=[("files", (tiff.name, f, "image/tiff"))],
                          data={"name": slide_id}).json()["id"]
    assert client.put(f"/api/analyses/{aid}/channels",
                      json={"mapping": [{"index": 0, "role": "red"}, {"index": 1, "role": "green"}]}).status_code == 200
    regs, objs = [], []
    for idx, (nr, ny) in fields.items():
        r, x0, y0 = field_region(idx)
        regs.append(r.model_dump(mode="json"))
        objs += [obj(f"r{idx}_{k}", x0 + 1 + k % 12, y0 + 2, ObjectClass.red) for k in range(nr)]
        objs += [obj(f"y{idx}_{k}", x0 + 1 + k % 12, y0 + 8, ObjectClass.yellow) for k in range(ny)]
    if not fields:   # a counting region that is not a lab field
        regs = [counting("c1", rect(5, 5, 100, 100)).model_dump(mode="json")]
    r = client.put(f"/api/analyses/{aid}/regions", json={"regions": regs})
    assert r.status_code == 200, r.text
    assert sorted(x["field_index"] for x in r.json()["regions"] if x["field_index"]) == sorted(fields)  # survives PUT
    detector.objects = objs
    j = client.post(f"/api/analyses/{aid}/detect", json={})
    assert j.status_code == 200, j.text
    import time
    for _ in range(500):
        s = client.get(f"/api/jobs/{j.json()['id']}").json()["status"]
        if s in ("done", "failed"):
            break
        time.sleep(0.02)
    assert s == "done"
    for rid in ([f"f{i}" for i in fields] or ["c1"]):
        act(client, aid, "region_inspected", region_id=rid, inspected=True)
        assert act(client, aid, "approve_region", region_id=rid).status_code == 200
    s = client.put(f"/api/analyses/{aid}/settings", json={"slide": meta(slide_id, series, **kw)})
    assert s.status_code == 200, s.text
    if finalize:
        assert client.post(f"/api/analyses/{aid}/finalize").status_code == 200
    return aid


def sheet(resp, name, data_only=False):
    return openpyxl.load_workbook(io.BytesIO(resp.content), data_only=data_only)[name]


def test_single_slide_export_names_content_and_formula(client, tiff, detector):
    # field -> (red-only count, yellow count); the first slide of the (synthetic) T7B template is the yellow series
    _, _, rows = read_template_csv(T7_CSV["T7B"])
    tpl_row = rows[0]
    t7b = [int(v) for v in tpl_row[9:27]]
    assert len(t7b) == 18 and all(v == "" for v in tpl_row[27:])
    aid = make(client, tiff, detector, {i + 1: (i % 3, n) for i, n in enumerate(t7b)}, slide_id=tpl_row[0])
    r = client.get(f"/api/analyses/{aid}/lab-export", params={"series": "T7B", "format": "xlsx"})
    assert r.status_code == 200, r.text
    assert 'filename="T7B_15-09-2026.xlsx"' in r.headers["content-disposition"] and r.headers["X-Export-Status"] == "FINAL"
    ws = sheet(r, "T7B")
    assert [c.value for c in ws[1]][:9] == ["slide ID", "SAMPLING DATE", "POLONY DATE ", "CONCENTRATION FACTOR",
                                           "SAMPLE VOLUM[ul]", "field [Pixel^2]", "grid [Pixel^2]", "QC",
                                           "FAGE ABUNDENCE [fage mL-1]"]
    assert ws["I2"].value == "=AVERAGE(J2:AG2)*(F2/G2)/E2*1000"
    assert [ws.cell(2, c).value for c in range(10, 28)] == t7b and ws["AB2"].value is None
    tpl_cached = openpyxl.load_workbook(TEMPLATES / T7_XLSX["T7B"], data_only=True)["T7B"]["I2"].value
    assert sheet(r, "T7B", True)["I2"].value == tpl_cached
    c = client.get(f"/api/analyses/{aid}/lab-export", params={"series": "T7B", "format": "csv"})
    assert 'filename="T7B_15-09-2026.csv"' in c.headers["content-disposition"]
    assert c.content.decode("utf-8-sig").split("\r\n")[1].startswith(
        ",".join([tpl_row[0], "15/09/2026", "29/09/2026", "1", "2.5", "2.9", "0.1", "1"] + tpl_row[8:]))
    # the other series maps to red-only: counts 0,1,2,0,1,2...
    a = client.get(f"/api/analyses/{aid}/lab-export", params={"series": "T7A", "format": "csv"})
    assert a.content.decode("utf-8-sig").split("\r\n")[1].split(",")[9:12] == ["0", "1", "2"]
    # series must be named when several are defined
    r = client.get(f"/api/analyses/{aid}/lab-export")
    assert r.status_code == 422 and r.json()["code"] == "series_required"
    assert client.get(f"/api/analyses/{aid}/lab-export", params={"series": "NOPE"}).status_code == 422
    # slide.csv / provenance carry the same lab abundance, series mapping and the field->region table
    sl = client.get(f"/api/analyses/{aid}/export", params={"kind": "slide"}).text.splitlines()
    row = dict(zip(sl[0].split(","), sl[1].split(",")))
    assert float(row["abundance_per_ml_T7B"]) == pytest.approx(float(tpl_row[8]), abs=1e-3)
    prov = client.get(f"/api/analyses/{aid}/export", params={"kind": "provenance"}).json()
    assert prov["lab"]["series"]["T7B"]["fields"][0] == {
        "field_index": 1, "region_id": "f1", "region_name": "field 1", "count": t7b[0],
        "counted_area_px2": prov["lab"]["series"]["T7B"]["fields"][0]["counted_area_px2"],
        "counted_area_fraction": 1.0, "note": None}
    assert {x["name"]: x["count_type"] for x in prov["lab"]["lab_series"]} == {"T7A": "red_only", "T7B": "yellow"}
    z = zipfile.ZipFile(io.BytesIO(client.get(f"/api/analyses/{aid}/export").content))
    assert "LAB FORMULA" in z.read("README.txt").decode()


def test_draft_and_incomplete_are_refused_with_exact_lists(client, tiff, detector):
    draft = make(client, tiff, detector, {1: (1, 2)}, finalize=False)
    r = client.get(f"/api/analyses/{draft}/lab-export", params={"series": "T7B"})
    assert r.status_code == 409 and r.json()["code"] == "not_finalized" and draft in r.json()["message"]
    r = client.post("/api/lab-export", json={"analysis_ids": [draft], "style": "counts"})
    assert r.status_code == 409
    incomplete = make(client, tiff, detector, {1: (1, 2)}, slide_id="ZZ_A", sampling_date=None, qc=None,
                      field_area=None, series=[("T7B", "yellow", None)])
    r = client.get(f"/api/analyses/{incomplete}/lab-export", params={"series": "T7B"})
    assert r.status_code == 422, r.text
    j = r.json()
    assert j["code"] == "lab_export_incomplete"
    assert j["detail"]["analyses"][0]["missing"] == ["sampling_date", "field_area", "qc"]
    assert "sampling_date, field_area, qc" in j["message"]
    nofield = make(client, tiff, detector, {}, slide_id="NF_A")
    j = client.get(f"/api/analyses/{nofield}/lab-export", params={"series": "T7B"}).json()
    assert j["detail"]["analyses"][0]["missing"][0].startswith("field_regions")


def test_batch_t7_zip_and_counts_workbook(client, tiff, detector):
    ids = [make(client, tiff, detector, {1: (1, 5), 2: (2, 7)}, slide_id="AI_A"),
           make(client, tiff, detector, {1: (3, 6), 2: (0, 8)}, slide_id="AI_B"),
           make(client, tiff, detector, {1: (0, 0), 2: (1, 1)}, slide_id="NTC_A"),
           make(client, tiff, detector, {1: (0, 0), 2: (0, 1)}, slide_id="NTC_B")]
    r = client.post("/api/lab-export", json={"analysis_ids": ids, "style": "t7", "format": "xlsx"})
    assert r.status_code == 200 and r.headers["content-type"] == "application/zip"
    z = zipfile.ZipFile(io.BytesIO(r.content))
    assert sorted(z.namelist()) == ["T7A_15-09-2026.xlsx", "T7B_15-09-2026.xlsx", "lab_export_notes.txt"]
    t7b = openpyxl.load_workbook(io.BytesIO(z.read("T7B_15-09-2026.xlsx")))["T7B"]
    assert [t7b.cell(r_, 1).value for r_ in range(2, 6)] == ["AI_A", "AI_B", "NTC_A", "NTC_B"]
    assert [t7b.cell(2, c).value for c in (10, 11)] == [5, 7] and [t7b.cell(3, c).value for c in (10, 11)] == [6, 8]
    t7a = openpyxl.load_workbook(io.BytesIO(z.read("T7A_15-09-2026.xlsx")))["T7A"]
    assert [t7a.cell(3, c).value for c in (10, 11)] == [3, 0]
    # a single series -> the file itself; csv; zip = xlsx + csv
    one = client.post("/api/lab-export", json={"analysis_ids": ids, "series": "T7A", "format": "csv"})
    assert one.status_code == 200 and 'filename="T7A_15-09-2026.csv"' in one.headers["content-disposition"]
    both = zipfile.ZipFile(io.BytesIO(client.post("/api/lab-export", json={"analysis_ids": ids, "series": "T7B",
                                                                           "format": "zip"}).content))
    assert sorted(both.namelist()) == ["T7B_15-09-2026.csv", "T7B_15-09-2026.xlsx", "lab_export_notes.txt"]
    # counts-style: one workbook, one sheet per series (sheet_title), NTC referenced, no #REF!
    c = client.post("/api/lab-export", json={"analysis_ids": ids, "style": "counts"})
    assert c.status_code == 200 and 'filename="Counts_29-09-2026.xlsx"' in c.headers["content-disposition"]
    wb = openpyxl.load_workbook(io.BytesIO(c.content))
    assert wb.sheetnames == ["Clade A", "Clade B", "Notes"]
    ws = wb["Clade B"]
    assert ws["AN7"].value == "=AM7-AM12" and ws["AQ9"].value == "=AH9-AM12" and ws["AN12"].value is None
    assert "#REF!" not in zipfile.ZipFile(io.BytesIO(c.content)).read("xl/worksheets/sheet2.xml").decode()
    assert client.post("/api/lab-export", json={"analysis_ids": ids, "style": "counts", "format": "csv"}).status_code == 422
    assert client.post("/api/lab-export", json={"analysis_ids": [], "style": "t7"}).status_code == 422
    assert client.post("/api/lab-export", json={"analysis_ids": ["0" * 12]}).status_code == 404


def test_settings_validation_for_lab_fields(client, tiff, detector):
    aid = make(client, tiff, detector, {1: (1, 1)}, finalize=False)
    put = lambda **kw: client.put(f"/api/analyses/{aid}/settings", json={"slide": meta("AI_A", SERIES, **kw)})  # noqa: E731
    assert put(field_area=0).status_code == 422
    assert put(grid_area=-1).status_code == 422
    assert put(qc=2).status_code == 422
    assert put(lab_series=[{"name": "T7A", "count_type": "yellow"}, {"name": "T7A", "count_type": "red_only"}]).status_code == 422
    assert put(lab_series=[{"name": "T7A", "count_type": "blue"}]).status_code == 422
    assert put().status_code == 200
