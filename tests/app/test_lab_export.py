"""Lab-format export (docs/app/LAB_EXPORT.md): T7-style table + Counts-style workbook, compared with
template files that have the layout of the researcher's own spreadsheets (synthetic data,
tests/test_data/lab_templates/make_templates.py) and hand-computed values."""
import io
import re
import zipfile

import openpyxl
import pytest

from fakes import counting, exclusion, rect
from labhelp import (COUNTS_XLSX, T7_CSV, T7_XLSX, TEMPLATES, evaluate, field_region, formula_cells, lab_state,
                     lo_recalc, read_template_csv, soffice, template_slides)
from polony_app.export import lab
from polony_app.schemas import LabCountType, ObjectClass

pytestmark = pytest.mark.skipif(not TEMPLATES.exists(), reason="lab template files not present")


def _counted(row):
    """Field counts of a template CSV row: fields 1..n counted, then blanks."""
    cells = row[9:33]
    n = sum(1 for v in cells if v != "")
    assert all(v != "" for v in cells[:n]), "first template rows are counted contiguously"
    return [int(v) for v in cells[:n]]


# The first sample pair of the T7B template (the same slides are rows 6-9 of the Counts template's "Clade B").
if TEMPLATES.exists():
    _T7B_RAW, _, _T7B_ROWS = read_template_csv(T7_CSV["T7B"])
    ROW_A, ROW_B = _T7B_ROWS[0], _T7B_ROWS[1]
    ID_A, ID_B = ROW_A[0], ROW_B[0]                     # "<sample>_A" / "<sample>_B"
    SAMPLE = ID_A.rsplit("_", 1)[0]
    FA, FB = _counted(ROW_A), _counted(ROW_B)          # 18 and 17 counted fields
    I2_CACHED = openpyxl.load_workbook(TEMPLATES / T7_XLSX["T7B"], data_only=True)["T7B"]["I2"].value
else:  # pragma: no cover - module is skipped
    ROW_A = ROW_B = []
    ID_A, ID_B, SAMPLE, FA, FB, I2_CACHED = "X_A", "X_B", "X", [1], [1], None


def padded(f):
    return f + [None] * (24 - len(f))


def an(*a, **k):
    return lab_state(*a, **k).analysis


def wb_of(data, **kw):
    return openpyxl.load_workbook(io.BytesIO(data), **kw)


# ============================================================== T7 table: numbers
def test_reproduces_template_number_for_first_slide():
    a = an(FA, slide_id=ID_A, series=(("T7B", LabCountType.yellow),))
    rows = lab.t7_table([a], "T7B")
    assert len(rows) == 1 and len(rows[0]) == 33
    r = rows[0]
    assert r[0] == ID_A and r[3:8] == [1, 2.5, 2.9, 0.1, 1]
    assert len(FA) == 18 and r[9:27] == FA and r[27:] == [None] * 6
    assert r[8] == pytest.approx(float(ROW_A[8]), abs=5e-5)   # the template CSV's 10-significant-digit value
    assert r[8] == I2_CACHED                                  # == the value cached in the template workbook
    # hand computation: mean * (field/grid) / volume * 1000
    assert sum(FA) / 18 * (2.9 / 0.1) / 2.5 * 1000 == r[8]
    # the concentration factor is NOT part of the lab formula
    b = an(FA, slide_id=ID_A, cf=7, series=(("T7B", LabCountType.yellow),))
    assert lab.t7_table([b], "T7B")[0][8] == r[8]


def test_counts_style_AK_with_volume_2_5056_reproduces_template():
    tpl = openpyxl.load_workbook(TEMPLATES / COUNTS_XLSX, data_only=True)["Clade B"]   # cached template values
    a = an(FA, slide_id=ID_A, volume=2.5056)
    b = an(FB, slide_id=ID_B, volume=2.5056)
    wb = wb_of(lab.write_counts_xlsx([a, b]), data_only=True)
    ws = wb["T7B"]
    assert ws["AK7"].value == pytest.approx(tpl["AK7"].value, rel=1e-12)
    assert ws["AK9"].value == pytest.approx(tpl["AK9"].value, rel=1e-12)
    assert ws["AK7"].value == pytest.approx(sum(FA) / len(FA) * 29 * 1000 / 2.5056, rel=1e-12)   # by hand
    assert ws["AL7"].value == pytest.approx(tpl["AL7"].value)
    assert ws["AM7"].value == pytest.approx(tpl["AM7"].value)
    assert ws["F6"].value == tpl["F6"].value == sum(FA) and ws["F8"].value == tpl["F8"].value == sum(FB)   # SUM
    assert ws["G7"].value == 29 and ws["AG7"].value == pytest.approx(tpl["AG7"].value)


def test_blank_vs_zero_is_kept_in_table_csv_and_xlsx():
    a = an({1: 0, 3: 2}, slide_id="Z_A", series=(("S", LabCountType.yellow),))      # field 1 = 0 objects, 2 = no region
    r = lab.t7_table([a], "S")[0]
    assert r[9:12] == [0, None, 2] and r[8] == pytest.approx(11600)    # mean(0, 2) = 1 (blank is not averaged)
    line = lab.write_t7_csv([a], "S").decode("utf-8-sig").split("\r\n")[1]
    assert re.search(r",0,,2,{22}$", line) or line.split(",")[9:12] == ["0", "", "2"]
    ws = wb_of(lab.write_t7_xlsx([a], "S"))["S"]
    assert ws["J2"].value == 0 and ws["K2"].value is None and ws["L2"].value == 2
    cw = wb_of(lab.write_counts_xlsx([a]))["S"]
    assert cw["H7"].value == 0 and cw["I7"].value is None and cw["J7"].value == 2
    assert evaluate(ws, ws["I2"].value) == r[8]


# ============================================================== T7: template structure
@pytest.mark.parametrize("name,series,cls", [("T7A", "T7A", ObjectClass.red), ("T7B", "T7B", ObjectClass.yellow)])
def test_csv_is_byte_identical_to_template(name, series, cls):
    raw, _, _ = read_template_csv(T7_CSV[name])
    analyses = template_slides(T7_CSV[name], series, cls)
    out = lab.write_t7_csv(analyses, series)
    assert out == raw, "BOM + CRLF + no trailing newline + 10-significant-digit abundance + dd/mm/yyyy"
    assert out.startswith(b"\xef\xbb\xbf") and b"\r\n" in out and not out.endswith(b"\n")


def test_csv_conventions():
    out = lab.write_t7_csv([an(FA, slide_id=ID_A)], "T7B")
    assert out.startswith(b"\xef\xbb\xbfslide ID,SAMPLING DATE,POLONY DATE ,CONCENTRATION FACTOR,SAMPLE VOLUM[ul],")
    text = out.decode("utf-8-sig")
    assert text.split("\r\n")[1].startswith(f"{ID_A},15/09/2026,29/09/2026,1,2.5,2.9,0.1,1,{ROW_A[8]},{FA[0]},{FA[1]},")
    assert text.split("\r\n")[1].endswith(f",{FA[-1]},,,,,,")
    # with the template's dates the line is the template's first data line
    same = lab.write_t7_csv([an(FA, slide_id=ID_A, sampling=ROW_A[1], polony=ROW_A[2])], "T7B")
    assert same.decode("utf-8-sig").split("\r\n")[1] == _T7B_RAW.decode("utf-8-sig").split("\r\n")[1]


@pytest.mark.parametrize("name", ["T7A", "T7B"])
def test_xlsx_header_and_formulas_equal_template(name):
    tpl = openpyxl.load_workbook(TEMPLATES / T7_XLSX[name])
    tws = tpl.active
    analyses = template_slides(T7_CSV[name], name, ObjectClass.yellow)
    ws = wb_of(lab.write_t7_xlsx(analyses, name))[name]
    assert [c.value for c in ws[1]] == [c.value for c in tws[1]]            # exact header strings (incl. 'POLONY DATE ')
    assert ws.max_row == tws.max_row == 25 and ws.max_column == tws.max_column == 33
    for r in range(2, 26):
        assert ws.cell(r, 9).value == tws.cell(r, 9).value                   # same formula text, same rows
        assert [c.value for c in ws[r]][:8] == [c.value for c in tws[r]][:8]  # id, dates, factors, qc
        assert [c.value for c in ws[r]][9:] == [c.value for c in tws[r]][9:]  # blank stays blank, 0 stays 0
        assert ws.cell(r, 2).number_format == tws.cell(r, 2).number_format == "mm-dd-yy"
    for col, w in {"A": 8.0, "B": 13.33, "C": 11.83, "D": 20.33, "E": 15.66, "F": 11.33, "I": 23.5}.items():
        assert round(ws.column_dimensions[col].width, 2) == round(tws.column_dimensions[col].width, 2) == w


@pytest.mark.parametrize("name", ["T7A", "T7B"])
def test_xlsx_formula_values_match_template_cache_and_csv(name):
    """Recompute every column-I formula with the tiny evaluator and compare with (a) the exporter's cached
    values, (b) the template's own cached values and (c) the template CSV abundance column."""
    tcache = openpyxl.load_workbook(TEMPLATES / T7_XLSX[name], data_only=True).active
    _, _, csvrows = read_template_csv(T7_CSV[name])
    analyses = template_slides(T7_CSV[name], name, ObjectClass.yellow)
    data = lab.write_t7_xlsx(analyses, name)
    ws, cached = wb_of(data)[name], wb_of(data, data_only=True)[name]
    for r in range(2, 26):
        v = evaluate(ws, ws.cell(r, 9).value)
        assert v == pytest.approx(tcache.cell(r, 9).value, rel=1e-12)
        assert cached.cell(r, 9).value == pytest.approx(v, rel=1e-12)
        assert float(csvrows[r - 2][8]) == pytest.approx(v, rel=1e-9)


@pytest.mark.skipif(not soffice(), reason="LibreOffice not installed")
def test_libreoffice_recalculation_agrees_with_cached_values():
    analyses = template_slides(T7_CSV["T7B"], "T7B", ObjectClass.yellow)
    data = lab.write_t7_xlsx(analyses, "T7B")
    mine, lo = wb_of(data, data_only=True)["T7B"], lo_recalc(data)["T7B"]
    for r in range(2, 26):
        assert lo.cell(r, 9).value == pytest.approx(mine.cell(r, 9).value, rel=1e-12)
    # and the Counts workbook
    a = an(FA, slide_id=ID_A, volume=2.5056)
    b = an(FB, slide_id=ID_B, volume=2.5056)
    cdata = lab.write_counts_xlsx([a, b])
    cm, cl = wb_of(cdata, data_only=True)["T7B"], lo_recalc(cdata)["T7B"]
    for coord in ("F6", "F8", "AG7", "AH7", "AK7", "AL7", "AM7", "AG9", "AH9", "AK9"):
        assert cl[coord].value == pytest.approx(cm[coord].value, rel=1e-12), coord


# ============================================================== Counts-style
def counts_ws(analyses, series="T7B"):
    data = lab.write_counts_xlsx(analyses)
    return wb_of(data)[series], wb_of(data, data_only=True)[series], data


def test_counts_layout_equals_template_cells():
    tpl = openpyxl.load_workbook(TEMPLATES / COUNTS_XLSX)["Clade B"]
    a = an(FA, slide_id=ID_A, volume=2.5056)
    b = an(FB, slide_id=ID_B, volume=2.5056)
    ws, _, _ = counts_ws([a, b])
    for coord in ["B4", "C4", "D4", "E4", "F4", "G4", "H4", "AG4", "AH4", "AI4", "AJ4", "AK4", "AL4", "AM4", "AN4",
                  "AO4", "AP4", "AQ4", "A4"]:
        assert ws[coord].value == tpl[coord].value, coord             # exact header strings (yellow series)
    for coord in ["F6", "F8", "AG7", "AH7", "AI7", "AK7", "AL7", "AM7", "G7", "G9", "AG9", "AH9", "AK9", "D6", "D8"]:
        assert ws[coord].value == tpl[coord].value, coord             # same formulas / constants / replicate letters
    assert ws["AJ7"].value == 2.5056 and tpl["AJ7"].value == "=5.4*11.6/25"   # documented difference: number, not formula
    for row, f in ((7, FA), (9, FB)):                                # same counts as the template, blanks kept
        assert [ws.cell(row, c).value for c in range(8, 32)] == [tpl.cell(row, c).value for c in range(8, 32)] \
            == padded(f)
    assert [ws.cell(7, c).value for c in range(8, 32)][:18] == FA
    assert [ws.cell(9, c).value for c in range(8, 32)][:17] == FB and ws["Y9"].value is None
    assert ws["C6"].value == "15-09-2026" and ws["C8"].value == SAMPLE and ws["B8"].value == "T7B"
    for rng in ("B6:B7", "C6:C7", "D6:D7", "E6:E7", "F6:F7", "B8:B9", "F8:F9", "AG4:AG5", "H4:AE5"):
        assert rng in {str(m) for m in ws.merged_cells.ranges}, rng
    assert ws["AK7"].number_format == "0.0E+00"
    # B-replicate row carries no pair statistics (template put #REF! there)
    assert ws["AL9"].value is None and ws["AM9"].value is None


def test_counts_pairs_by_slide_id_and_formulas_evaluate_to_cached_values():
    ans = [an(FA, slide_id=ID_A, volume=2.5056), an(FB, slide_id=ID_B, volume=2.5056),
           an([2, 1, 3, 2, 9], slide_id="P2_A", volume=2.5056), an([5, 4, 4], slide_id="P2_B", volume=2.5056),
           an([4, 4], slide_id="SOLO", volume=2.5056)]
    ws, cached, _ = counts_ws(ans)
    assert [ws[f"D{r}"].value for r in (6, 8, 11, 13, 16)] == ["A", "B", "A", "B", None]
    assert [ws[f"C{r}"].value for r in (8, 13, 16)] == [SAMPLE, "P2", "SOLO"]
    assert ws["AL12"].value == "=100*ABS(AH12-AH14)/AVERAGE(AH12,AH14)" and ws["AL17"].value is None
    for coord, f in formula_cells(ws).items():
        assert cached[coord].value == pytest.approx(evaluate(ws, f), rel=1e-12), (coord, f)


def test_counts_ntc_blank_without_ntc_and_references_with_ntc():
    base = [an([10, 12], slide_id="S1_A", volume=2.5), an([11, 13], slide_id="S1_B", volume=2.5)]
    ws, cached, data = counts_ws(base, "T7B")
    assert ws["AN7"].value is None and ws["AQ7"].value is None and ws["AQ9"].value is None
    assert b"#REF!" not in zipfile.ZipFile(io.BytesIO(data)).read("xl/worksheets/sheet1.xml")
    ntc = [an([1, 0], slide_id="NTC_A", volume=2.5), an([0, 1], slide_id="NTC_B", volume=2.5)]
    ws, cached, data = counts_ws(base + ntc, "T7B")
    assert ws["AN7"].value == "=AM7-AM12" and ws["AQ7"].value == "=AH7-AM12" and ws["AQ9"].value == "=AH9-AM12"
    assert ws["AN12"].value is None and ws["AQ12"].value is None                  # NTC rows themselves
    assert cached["AN7"].value == pytest.approx(evaluate(ws, ws["AN7"].value))
    assert cached["AQ9"].value == pytest.approx(evaluate(ws, ws["AQ9"].value))
    assert "#REF!" not in b"".join(zipfile.ZipFile(io.BytesIO(data)).read(n) for n in zipfile.ZipFile(io.BytesIO(data)).namelist()).decode("utf-8", "ignore")


def test_counts_efficiency_only_if_input_set_and_dilution_applied():
    a = an([10, 20], slide_id="E_A", volume=2.5, cf=4, efficiency=1500)
    b = an([10, 20], slide_id="E_B", volume=2.5, cf=4, efficiency=1500)
    ws, cached, _ = counts_ws([a, b])
    assert ws["AO7"].value == "=100*AM7/1500" and ws["AI7"].value == 4
    assert ws["AK7"].value == "=(AH7)*1000/AJ7*AI7"                      # dilution IS applied in the Counts layout
    assert cached["AK7"].value == pytest.approx(15 * 29 * 1000 / 2.5 * 4)
    ws2, _, _ = counts_ws([an([10], slide_id="F_A"), an([10], slide_id="F_B")])
    assert ws2["AO7"].value is None


def test_counts_refuses_ambiguous_replicates_and_csv():
    with pytest.raises(lab.LabExportError) as e:
        lab.write_counts_xlsx([an([1], slide_id="X_A", aid="a" * 12), an([2], slide_id="X", aid="b" * 12)])
    assert e.value.code == "ambiguous_replicates"
    with pytest.raises(lab.LabExportError) as e:
        lab.build_files([an([1])], "counts", None, "csv")
    assert e.value.status == 422


# ============================================================== Notes / warnings
def test_notes_sheet_lists_mapping_ids_hashes_warnings():
    ex = exclusion("xx", rect(12, 5, 20, 20))                      # covers the right half of field 1 ((5..20)x(5..20))
    ex2 = exclusion("yy", rect(23, 5, 38, 20))                     # covers field 2 entirely
    a = an({1: 3, 2: 4, 3: 5}, slide_id="N_A", extra_regions=[ex, ex2], series=(("T7B", LabCountType.yellow),),
           aid="c" * 12)
    r = lab.t7_table([a], "T7B")[0]
    assert r[9:12] == [3, None, 5]            # fully excluded field: BLANK, not a counted zero
    ws = wb_of(lab.write_t7_xlsx([a], "T7B"))
    assert ws.sheetnames == ["T7B", "Notes"]
    text = "\n".join(" | ".join(str(c.value) for c in row if c.value is not None) for row in ws["Notes"].iter_rows())
    for needle in ("partly covered by exclusion", "entirely covered by exclusion", "T7B", "yellow", "N_A",
                   "cccccccccccc", "FINAL", "0000000000000000000000000000000000000000000000000000000000000000",
                   "polony_app 0.1.0", "AVERAGE(fields)", "field 1"):
        assert needle in text, needle
    assert "only" in text and "NOT scaled" in text


def test_uncertain_and_non_field_regions_are_warned():
    from fakes import obj
    st = lab_state({1: 2}, extra_regions=[counting("loose", rect(100, 60, 140, 100))])
    st.analysis.objects.append(obj("u", 8, 8, ObjectClass.uncertain, region_id="f1"))
    st.analysis.objects[-1].region_id = "f1"
    slides = lab.prepare([st.analysis], "T7B")
    w = " ".join(slides[0].warnings)
    assert "without a field_index" in w and "loose" in w and "uncertain" in w


# ============================================================== refusals
def test_refusals_unfinalized_missing_metadata_field_index():
    draft = an(FA, slide_id=ID_A, finalized=False, aid="d" * 12)
    with pytest.raises(lab.LabExportError) as e:
        lab.write_t7_csv([draft], "T7B")
    assert e.value.status == 409 and e.value.code == "not_finalized" and ID_A in e.value.message

    bare = lab_state({}, slide_id=None, sampling=None, polony=None, volume=None, field_area=None, grid_area=None,
                     cf=None, qc=None, series=(), confirmed=False).analysis
    bare.slide.slide_id = None
    with pytest.raises(lab.LabExportError) as e:
        lab.t7_table([bare], "T7B")
    assert e.value.status == 422 and e.value.code == "lab_export_incomplete"
    missing = e.value.detail["analyses"][0]["missing"]
    assert missing == ["slide_id", "sampling_date", "polony_date", "sample_volume_ul", "field_area", "grid_area",
                       "concentration_factor", "qc",
                       "concentration_formula_confirmed (field geometry and series mapping not confirmed)",
                       "lab_series"]
    for m in ("slide_id", "sampling_date", "qc", "lab_series"):
        assert m in e.value.message

    nofield = an({}, series=(("T7B", LabCountType.yellow),))
    with pytest.raises(lab.LabExportError) as e:
        lab.t7_table([nofield], "T7B")
    assert e.value.detail["analyses"][0]["missing"][0].startswith("field_regions")

    other = an(FA)
    with pytest.raises(lab.LabExportError) as e:
        lab.t7_table([other], "T7A")                      # series not defined for this slide
    assert "series 'T7A' is not defined" in e.value.message

    bad_date = an(FA, sampling="2026-09-15")
    with pytest.raises(lab.LabExportError) as e:
        lab.t7_table([bad_date], "T7B")
    assert "sampling_date (not dd/mm/yyyy" in e.value.message


def test_field_index_over_24_and_duplicates_refused():
    st = lab_state({1: 1})
    r25 = counting("f25", rect(100, 60, 120, 80), name="far")
    r25.field_index = 25
    st.analysis.regions.append(r25)
    with pytest.raises(lab.LabExportError) as e:
        lab.t7_table([st.analysis], "T7B")
    assert e.value.status == 422 and "field_index 25" in e.value.message and "24" in e.value.message

    st = lab_state({1: 1})
    dup = counting("dup", rect(100, 60, 120, 80), name="dupfield")
    dup.field_index = 1
    st.analysis.regions.append(dup)
    with pytest.raises(lab.LabExportError) as e:
        lab.t7_table([st.analysis], "T7B")
    assert "duplicate field_index 1" in e.value.message


def test_inconsistent_series_mapping_and_duplicate_ids_refused():
    a = an([1], slide_id="Q_A", aid="1" * 12, series=(("T7B", LabCountType.yellow),))
    b = an([1], slide_id="Q_B", aid="2" * 12, series=(("T7B", LabCountType.red_only),))
    with pytest.raises(lab.LabExportError) as e:
        lab.t7_table([a, b], "T7B")
    assert e.value.code == "inconsistent_series"
    c = an([1], slide_id="Q_A", aid="3" * 12)
    with pytest.raises(lab.LabExportError) as e:
        lab.t7_table([a, c], "T7B")
    assert e.value.code == "duplicate_slide_id"


def test_count_types_select_the_right_counts():
    from fakes import obj
    st = lab_state({1: 0}, series=(("R", LabCountType.red_only), ("G", LabCountType.green_only), ("Y", LabCountType.yellow),
                                   ("RP", LabCountType.red_positive), ("GP", LabCountType.green_positive)))
    a = st.analysis
    for i, c in enumerate([ObjectClass.red] * 2 + [ObjectClass.green] * 3 + [ObjectClass.yellow] * 4):
        a.objects.append(obj(f"m{i}", 6 + i, 8, c, region_id="f1"))
    got = {s: lab.t7_table([a], s)[0][9] for s in ("R", "G", "Y", "RP", "GP")}
    assert got == {"R": 2, "G": 3, "Y": 4, "RP": 6, "GP": 7}
