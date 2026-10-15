"""Generate the SYNTHETIC lab-format templates used by tests/app/test_lab_export.py.

The files in this folder mimic the layout of the laboratory's own spreadsheets (the T7-style
count tables and the Counts-style workbook described in docs/app/LAB_EXPORT.md), but every
slide id, date and count in them is made up:

* ``T7A_template.csv`` / ``T7A_template.xlsx`` - red-only style series (small counts, many zeros)
* ``T7B_template.csv`` / ``T7B_template.xlsx`` - yellow style series (counts ~5-40)
* ``Counts_template.xlsx`` - sheets "Clade B" (= the T7B slides), "Clade A" (= the T7A slides)
  and "Controls" (a positive control and an NTC, twice)

24 slides per series: samples S01..S12, replicates _A/_B. Each slide has a random number of
counted fields (6-24, field 1.. in order); a blank cell means "field not counted", 0 means
"counted, zero objects" - both occur. A few rows are pinned so the tests can rely on them
(see ``_pinned``).

Layout details reproduced from the lab files (the tests compare against them):
* T7 header strings (incl. the trailing space in "POLONY DATE " and the lab's spellings),
  25 rows x 33 columns, real Excel dates with number format ``mm-dd-yy``, column widths,
  column I formula ``=AVERAGE(Jn:AGn)*(Fn/Gn)/En*1000``.
* CSVs as Excel writes them: UTF-8 BOM, CRLF, no newline after the last row, dd/mm/yyyy dates,
  "General" numbers (no trailing .0), abundance with 10 significant digits, blank = empty.
* Counts workbook: row-4 headers, two rows per slide, merged B..F cells, ``=SUM``/``=AVERAGE``
  formulas, ``AJ = =5.4*11.6/25``, efficiency ``/1500`` and ``/1000`` columns, and the
  ``#REF!`` formulas the lab's workbook carries in the NTC columns (the app never writes those).

Every formula cell carries a CACHED value computed here in Python (openpyxl cannot write cached
values, so the saved package is post-processed to inject ``<v>`` elements; error results are
written as ``t="e"`` cells). ``openpyxl.load_workbook(..., data_only=True)`` reads them back.
No customXml parts; the core properties name no person; zip timestamps are fixed, so the output
is byte-reproducible.

Run (standard library + openpyxl only):  python tests/test_data/lab_templates/make_templates.py
"""
from __future__ import annotations

import csv
import io
import random
import re
import zipfile
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter

OUT = Path(__file__).resolve().parent
SEED = 20260303
SAMPLING = (3, 3, 2026)          # dd, mm, yyyy  (synthetic)
POLONY = (17, 3, 2026)
HYBRIDIZATION = (18, 3, 2026)
CONTROLS_SLIDES = (10, 2, 2026)
CREATOR = "Polony Review test data"
FIXED_TIME = datetime(2026, 3, 18, 12, 0, 0)
N_FIELDS = 24
SAMPLES = [f"S{i:02d}" for i in range(1, 13)]

T7_HEADER = (["slide ID", "SAMPLING DATE", "POLONY DATE ", "CONCENTRATION FACTOR", "SAMPLE VOLUM[ul]",
              "field [Pixel^2]", "grid [Pixel^2]", "QC", "FAGE ABUNDENCE [fage mL-1]"]
             + [f"field {i}" for i in range(1, N_FIELDS + 1)])
T7_WIDTHS = {"A": 8.0, "B": 13.33203125, "C": 11.83203125, "D": 20.33203125, "E": 15.6640625,
             "F": 11.33203125, "G": 11.0, "H": 11.0, "I": 23.5}
CF, VOLUME, FIELD_AREA, GRID_AREA = 1, 2.5, 2.9, 0.1
COUNTS_G = 29                    # field/grid as the Counts workbook writes it
COUNTS_AJ = "=5.4*11.6/25"       # sample volume per slide (formula in the lab's workbook)


def dslash(d):
    return f"{d[0]:02d}/{d[1]:02d}/{d[2]}"


def ddash(d):
    return f"{d[0]:02d}-{d[1]:02d}-{d[2]}"


# ------------------------------------------------------------------ synthetic slide rows
def _fields(rng, n, draw):
    return [draw() for _ in range(n)] + [None] * (N_FIELDS - n)


def _pinned(series, sid, f):
    """Rows the tests rely on (documented there)."""
    if series == "T7A" and sid == "S01_A":        # 18 counted cells, zeros inside, then 6 blanks
        f[:3] = [1, 3, 0]
        f[3] = 0
        f[17] = 0
    if series == "T7B" and sid == "S03_A":        # a counted zero in the yellow series
        f[1] = 0
    if sid == "S07_B":                            # a blank (not counted) field between counted ones
        f[4] = None
    return f


def slide_rows(series, rng):
    rows = []
    for k, s in enumerate(SAMPLES):
        base = 20 if k == 0 else rng.randint(5, 35)
        for rep in "AB":
            sid = f"{s}_{rep}"
            if sid == "S01_A":
                n = 18
            elif sid == "S01_B":
                n = 17
            elif s == "S11":
                n = 24                                # a fully counted slide
            else:
                n = rng.randint(6, 24)
            if series == "T7B":
                f = _fields(rng, n, lambda: max(0, base + rng.randint(-8, 8)))
            else:
                f = _fields(rng, n, lambda: rng.choice([0, 0, 0, 0, 1, 1, 2, 3]))
            if not any(f):                        # never an all-zero slide (pair statistics divide by it)
                f[0] = 1
            f = _pinned(series, sid, f)
            qc = 0 if (series == "T7B" and s == "S04") else 1
            rows.append({"id": sid, "qc": qc, "fields": f})
    return rows


def abundance(fields):
    """Python mirror of =AVERAGE(J:AG)*(F/G)/E*1000 (same operation order as the Excel cell)."""
    v = [x for x in fields if x is not None]
    return sum(v) / len(v) * (FIELD_AREA / GRID_AREA) / VOLUME * 1000


# ------------------------------------------------------------------ CSV
def _general(v):
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v)


def t7_csv(rows) -> bytes:
    buf = io.StringIO()
    w = csv.writer(buf, lineterminator="\r\n")
    w.writerow(T7_HEADER)
    for r in rows:
        w.writerow([r["id"], dslash(SAMPLING), dslash(POLONY), _general(CF), _general(VOLUME), _general(FIELD_AREA),
                    _general(GRID_AREA), r["qc"], format(abundance(r["fields"]), ".10g")]
                   + ["" if x is None else x for x in r["fields"]])
    text = buf.getvalue()
    if text.endswith("\r\n"):
        text = text[:-2]
    return b"\xef\xbb\xbf" + text.encode("utf-8")


# ------------------------------------------------------------------ xlsx helpers
_STAMP = FIXED_TIME.strftime("%Y-%m-%dT%H:%M:%SZ")
CORE_XML = (
    '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n'
    '<cp:coreProperties xmlns:cp="http://schemas.openxmlformats.org/package/2006/metadata/core-properties" '
    'xmlns:dc="http://purl.org/dc/elements/1.1/" xmlns:dcterms="http://purl.org/dc/terms/" '
    'xmlns:dcmitype="http://purl.org/dc/dcmitype/" xmlns:xsi="http://www.w3.org/2001/XMLSchema-instance">'
    f'<dc:title>Synthetic lab-format template (test data)</dc:title><dc:creator>{CREATOR}</dc:creator>'
    f'<cp:lastModifiedBy>{CREATOR}</cp:lastModifiedBy>'
    f'<dcterms:created xsi:type="dcterms:W3CDTF">{_STAMP}</dcterms:created>'
    f'<dcterms:modified xsi:type="dcterms:W3CDTF">{_STAMP}</dcterms:modified></cp:coreProperties>')


def _new_workbook():
    wb = Workbook()
    p = wb.properties
    p.creator = p.lastModifiedBy = CREATOR      # (core.xml is replaced by CORE_XML on save)
    return wb


def _save(wb, cached: dict) -> bytes:
    """Save, inject the cached formula results, rewrite the zip with fixed timestamps."""
    buf = io.BytesIO()
    wb.save(buf)
    zin = zipfile.ZipFile(io.BytesIO(buf.getvalue()))
    sheet_files = {ws.title: f"xl/worksheets/sheet{i}.xml" for i, ws in enumerate(wb.worksheets, start=1)}
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            data = zin.read(item.filename)
            for title, fname in sheet_files.items():
                if item.filename == fname and cached.get(title):
                    xml = data.decode("utf-8")
                    for coord, val in cached[title].items():
                        pat = re.compile(r'<c r="%s"([^>]*)>(<f>[^<]*</f>)<v\s*/>' % coord)
                        if isinstance(val, str):          # an Excel error such as #REF!
                            rep = lambda m, v=val: f'<c r="{coord}"{m.group(1)} t="e">{m.group(2)}<v>{v}</v>'
                        else:
                            rep = lambda m, v=val: f'<c r="{coord}"{m.group(1)}>{m.group(2)}<v>{_v(v)}</v>'
                        xml, n = pat.subn(rep, xml, count=1)
                        assert n == 1, f"formula cell {coord} not found in {title}"
                    data = xml.encode("utf-8")
            if item.filename == "docProps/core.xml":    # openpyxl stamps "modified" with the current time
                data = CORE_XML.encode("utf-8")
            info = zipfile.ZipInfo(item.filename, date_time=FIXED_TIME.timetuple()[:6])
            info.compress_type = zipfile.ZIP_DEFLATED
            info.external_attr = 0o644 << 16
            zout.writestr(info, data)
    blob = out.getvalue()
    assert not any(n.lower().startswith("customxml") for n in zipfile.ZipFile(io.BytesIO(blob)).namelist())
    return blob


def _v(x):
    if isinstance(x, int) or (isinstance(x, float) and x.is_integer() and abs(x) < 1e15):
        return str(int(x))
    return repr(float(x))


# ------------------------------------------------------------------ T7 workbook
def t7_xlsx(series, rows) -> bytes:
    wb = _new_workbook()
    ws = wb.active
    ws.title = series
    ws.sheet_format.defaultColWidth = 8.83203125
    ws.sheet_format.baseColWidth = 10
    for k, w in T7_WIDTHS.items():
        ws.column_dimensions[k].width = w
    ws.append(T7_HEADER)
    cached = {}
    for i, r in enumerate(rows):
        n = i + 2
        ws.cell(n, 1, r["id"])
        for c, d in ((2, SAMPLING), (3, POLONY)):
            cell = ws.cell(n, c, datetime(d[2], d[1], d[0]))
            cell.number_format = "mm-dd-yy"
        for c, v in ((4, CF), (5, VOLUME), (6, FIELD_AREA), (7, GRID_AREA), (8, r["qc"])):
            ws.cell(n, c, v)
        ws.cell(n, 9, f"=AVERAGE(J{n}:AG{n})*(F{n}/G{n})/E{n}*1000")
        cached[f"I{n}"] = abundance(r["fields"])
        for j, x in enumerate(r["fields"]):
            if x is not None:
                ws.cell(n, 10 + j, x)
    return _save(wb, {series: cached})


# ------------------------------------------------------------------ Counts workbook
REF = "#REF!"
HEADS = {
    "A4": "Cy5 red & Cy3 green \n(yellow counts)", "B4": "assay", "C4": "Phage & dilution",
    "D4": "Slide number", "E4": "yellow", "F4": "summ / slide", "G4": "well area/grid area",
    "H4": "total yellow polonies/ grid", "AG4": "average number of polonies/grid",
    "AH4": "polonies/ slide", "AI4": "dilution", "AJ4": "sample volume/ slide", "AK4": " polonies/ ml",
    "AL4": "%v diff btw slides", "AM4": "Average ", "AN4": "Average - NTC", "AO4": "Efficiency average",
    "AP4": "Efficiency average-NTC", "AQ4": "pol/slide -NTC",
}


def _avg(*xs):
    if any(x == REF for x in xs):
        return REF
    v = [x for x in xs if x is not None]
    return sum(v) / len(v)


def _counts_frame(ws, title, subtitle, slides_from):
    ws.sheet_format.defaultRowHeight = 16
    ws.sheet_format.baseColWidth = 10
    for k, w in {"A": 18.6640625, "B": 13.0, "C": 14.5, "D": 11.6640625, "E": 17.1640625, "F": 14.0}.items():
        ws.column_dimensions[k].width = w
    ws.column_dimensions["H"].width = 0
    ws.column_dimensions["H"].hidden = True
    for c in range(9, 33):
        ws.column_dimensions[get_column_letter(c)].width = 10.83203125
        ws.column_dimensions[get_column_letter(c)].hidden = True
    ws["A1"], ws["B1"] = title, subtitle
    ws["A3"] = f"Slides from {ddash(slides_from)}"
    for k in ("A1", "B1", "A3"):
        ws[k].font = Font(size=16)
    for k, v in HEADS.items():
        ws[k] = v
        ws[k].alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for k in ["B", "C", "D", "E", "F", "G", "AG", "AH", "AI", "AJ", "AK"]:
        ws.merge_cells(f"{k}4:{k}5")
    ws.merge_cells("H4:AE5")


def _slide_block(ws, cache, h, *, label_a, label_b, assay, rep, fields, e_value=None):
    """Header row h + value row v=h+1 of one slide (formulas cached in `cache`)."""
    v = h + 1
    if rep == "A":
        ws[f"C{h}"] = label_a
    else:
        ws[f"B{h}"] = assay
        if label_b is not None:
            ws[f"C{h}"] = label_b
    ws[f"D{h}"] = rep
    if e_value is not None:
        ws[f"E{h}"] = e_value
    for col in "BCDEF":
        ws.merge_cells(f"{col}{h}:{col}{v}")
    ws[f"F{h}"] = f"=SUM(H{v}:AE{v})"
    counted = [x for x in fields if x is not None]
    cache[f"F{h}"] = sum(counted)
    for i in range(N_FIELDS):
        ws.cell(h, 8 + i, f"Field-{i + 1}")
        if fields[i] is not None:
            ws.cell(v, 8 + i, fields[i])
    ws[f"G{v}"] = COUNTS_G
    ws[f"G{v}"].number_format = "0.00"
    ag = sum(counted) / len(counted)
    ah = ag * COUNTS_G
    aj = 5.4 * 11.6 / 25
    ws[f"AG{v}"], cache[f"AG{v}"] = f"=AVERAGE(H{v}:AE{v})", ag
    ws[f"AH{v}"], cache[f"AH{v}"] = f"=AG{v}*G{v}", ah
    ws[f"AI{v}"] = CF
    ws[f"AJ{v}"], cache[f"AJ{v}"] = COUNTS_AJ, aj
    ws[f"AK{v}"], cache[f"AK{v}"] = f"=(AH{v})*1000/AJ{v}*AI{v}", (ah) * 1000 / aj * CF
    for col, nf in {"AG": "0.00", "AH": "0.0", "AK": "0.0E+00"}.items():
        ws[f"{col}{v}"].number_format = nf
    return v


def _pair_stats(ws, cache, va, vb, *, b_partner=None):
    """Replicate statistics as in the lab workbook (incl. its #REF! NTC references)."""
    aha, ahb = cache[f"AH{va}"], cache[f"AH{vb}"]
    ws[f"AL{va}"] = f"=100*ABS(AH{va}-AH{vb})/AVERAGE(AH{va},AH{vb})"
    cache[f"AL{va}"] = 100 * abs(aha - ahb) / _avg(aha, ahb)
    ws[f"AM{va}"], cache[f"AM{va}"] = f"=AVERAGE(AH{va},AH{vb})", _avg(aha, ahb)
    if b_partner is None:
        ws[f"AM{vb}"], cache[f"AM{vb}"] = f"=AVERAGE(AH{vb},{REF})", REF
    else:   # the Controls sheet averages the B row with an (empty) header row below it
        ws[f"AM{vb}"], cache[f"AM{vb}"] = f"=AVERAGE(AH{vb},AH{b_partner})", _avg(ahb, None)
    for v in (va, vb):
        _ntc_and_efficiency(ws, cache, v)


def _ntc_and_efficiency(ws, cache, v):
    ws[f"AN{v}"], cache[f"AN{v}"] = f"=AM{v}-{REF}", REF
    am = cache[f"AM{v}"]
    ws[f"AO{v}"], cache[f"AO{v}"] = f"=100*AM{v}/1500", (REF if am == REF else 100 * am / 1500)
    ws[f"AP{v}"], cache[f"AP{v}"] = f"=100*AN{v}/1000", REF
    ws[f"AQ{v}"], cache[f"AQ{v}"] = f"=AH{v}-{REF}", REF
    for col, nf in (("AL", "0.00"), ("AM", "0.0"), ("AN", "0.00"), ("AO", "0.00"), ("AQ", "0.0")):
        ws[f"{col}{v}"].number_format = nf


def counts_xlsx(t7b_rows, t7a_rows, rng) -> bytes:
    wb = _new_workbook()
    wb.remove(wb.active)
    cached = {}
    for sheet, rows in (("Clade B", t7b_rows), ("Clade A", t7a_rows)):
        ws = wb.create_sheet(sheet)
        cache = cached.setdefault(sheet, {})
        _counts_frame(ws, "Synthetic assay", f"{sheet} counts", POLONY)
        by_id = {r["id"]: r for r in rows}
        h = 6
        for s in SAMPLES:
            a, b = by_id[f"{s}_A"], by_id[f"{s}_B"]
            va = _slide_block(ws, cache, h, label_a=ddash(SAMPLING), label_b=s, assay=sheet, rep="A",
                              fields=a["fields"], e_value=rng.choice([None, None, 0, 1, 2]))
            vb = _slide_block(ws, cache, h + 2, label_a=None, label_b=s, assay=sheet, rep="B", fields=b["fields"])
            _pair_stats(ws, cache, va, vb)
            h += 5
    ws = wb.create_sheet("Controls")
    cache = cached.setdefault("Controls", {})
    _counts_frame(ws, "Controls (synthetic)", "Controls counts", CONTROLS_SLIDES)
    for h in (6, 16):
        pos = [_fields(rng, 12, lambda: rng.randint(10, 30)) for _ in range(2)]
        va = _slide_block(ws, cache, h, label_a="PC", label_b=None, assay="Controls", rep="A", fields=pos[0], e_value=0)
        vb = _slide_block(ws, cache, h + 2, label_a=None, label_b=None, assay="Controls", rep="B", fields=pos[1])
        _pair_stats(ws, cache, va, vb, b_partner=h + 5)
        ntc = _fields(rng, 12, lambda: rng.choice([0, 0, 0, 1]))
        ntc[0] = 1                                   # never an all-zero NTC
        vn = _slide_block(ws, cache, h + 5, label_a="NTC", label_b=None, assay="Controls", rep="A", fields=ntc)
        ws[f"AL{vn}"], cache[f"AL{vn}"] = f"=100*ABS(AH{vn}-{REF})/AVERAGE(AH{vn},{REF})", REF
        ws[f"AM{vn}"], cache[f"AM{vn}"] = f"=AVERAGE(AH{vn},{REF})", REF
        _ntc_and_efficiency(ws, cache, vn)
    return _save(wb, cached)


# ------------------------------------------------------------------ main
def build() -> dict[str, bytes]:
    rng = random.Random(SEED)
    t7a = slide_rows("T7A", rng)
    t7b = slide_rows("T7B", rng)
    return {
        "T7A_template.csv": t7_csv(t7a),
        "T7B_template.csv": t7_csv(t7b),
        "T7A_template.xlsx": t7_xlsx("T7A", t7a),
        "T7B_template.xlsx": t7_xlsx("T7B", t7b),
        "Counts_template.xlsx": counts_xlsx(t7b, t7a, rng),
    }


def main():
    for name, data in build().items():
        (OUT / name).write_bytes(data)
        print(f"wrote {name} ({len(data)} bytes)")


if __name__ == "__main__":
    main()
